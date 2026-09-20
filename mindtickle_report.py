"""
Mindtickle voortgangsrapport generator
---------------------------------------
Haalt automatisch de leerling-voortgang op uit Mindtickle en zet dit om
in een opgemaakt Excel-bestand, klaar om te versturen naar de groepsapp.

Gebruik:
    1. Vul .env in (kopieer .env.example) met je cookie en manager ID.
    2. pip install -r requirements.txt
    3. python mindtickle_report.py

Het script:
    - Haalt alle leerlingen op via de Mindtickle "members" API
    - Verwijdert e-mailadressen uit het Excel-overzicht
    - Houdt bij dubbele namen alleen de hoogste voltooiing bij
    - Berekent per persoon of ze boven de score-drempel zitten (standaard 50%)
    - Sorteert op score (hoogste bovenaan)
    - Kleurt rood/groen o.b.v. de drempel
    - Slaat op als voortgang_overzicht_<datum>.xlsx
"""

import os
import sys
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, urlparse
import json

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

load_dotenv()

# ── Configuratie (uit .env) ──────────────────────────────────────────────
MINDTICKLE_SUBDOMAIN = os.getenv("MINDTICKLE_SUBDOMAIN", "directresultmarketing")
MANAGER_ID = os.getenv("MINDTICKLE_MANAGER_ID")
COOKIE = os.getenv("MINDTICKLE_COOKIE")
HAR_PATH = os.getenv("MINDTICKLE_HAR")
SCORE_THRESHOLD = float(os.getenv("SCORE_THRESHOLD", "50"))  # in procenten


def load_har_request(har_path):
    """Leest endpoint, manager-ID en cookie uit een geëxporteerde HAR."""
    try:
        with Path(har_path).open("r", encoding="utf-8") as har_file:
            har = json.load(har_file)
    except (OSError, json.JSONDecodeError) as exc:
        sys.exit(f"❌ HAR-bestand kon niet worden gelezen: {exc}")

    for entry in har.get("log", {}).get("entries", []):
        request = entry.get("request", {})
        url = request.get("url", "")
        if "/api/node/manager/" not in url or not urlparse(url).path.endswith("/members"):
            continue

        parsed_url = urlparse(url)
        manager_part = parsed_url.path.split("/manager/", 1)[-1].split("/", 1)[0]
        header_cookie = next(
            (header.get("value") for header in request.get("headers", [])
             if header.get("name", "").lower() == "cookie"),
            "",
        )
        cookie = header_cookie or "; ".join(
            f"{item.get('name')}={item.get('value')}"
            for item in request.get("cookies", [])
            if item.get("name")
        )
        return {
            "subdomain": parsed_url.hostname.split(".", 1)[0],
            "manager_id": manager_part,
            "cookie": cookie,
            "query": parse_qs(parsed_url.query),
        }

    sys.exit("❌ Geen Mindtickle members-request gevonden in het HAR-bestand.")


def normalize_cookie(cookie: str) -> str:
    """Verwijdert regeleinden die ontstaan bij het plakken van een cookie."""
    return " ".join(line.strip() for line in cookie.splitlines() if line.strip()).strip()


def fetch_members():
    """Haalt alle leerlingen in één keer op via de members-API."""
    subdomain = MINDTICKLE_SUBDOMAIN
    manager_id = MANAGER_ID
    cookie = COOKIE
    har_query = {}

    if HAR_PATH:
        har_config = load_har_request(HAR_PATH)
        subdomain = subdomain or har_config["subdomain"]
        manager_id = manager_id or har_config["manager_id"]
        cookie = cookie or har_config["cookie"]
        har_query = har_config["query"]

    cookie = normalize_cookie(cookie or "")
    if not manager_id or not cookie:
        sys.exit(
            "❌ Mindtickle-authenticatie ontbreekt.\n"
            "   Vul MINDTICKLE_COOKIE in, of gebruik MINDTICKLE_HAR met een HAR "
            "waarin de cookie staat."
        )

    base_url = f"https://{subdomain}.mindtickle.com/api/node/manager/{manager_id}/members"

    headers = {
        "accept": "application/json",
        "cookie": cookie,
        "user-agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36"
        ),
    }

    # Stap 1: klein verzoek om het totaal aantal leerlingen te weten
    probe_params = {
        "invitedOnAfter": har_query.get("invitedOnAfter", ["all-time"])[0],
        "pagesize": 1,
        "sort": "learnerName",
    }
    r = requests.get(base_url, headers=headers, params=probe_params, timeout=30)
    if r.status_code == 401 or r.status_code == 403:
        sys.exit(
            "❌ Niet geautoriseerd (401/403). Je cookie is waarschijnlijk verlopen.\n"
            "   Log opnieuw in op Mindtickle en kopieer een verse cookie naar .env."
        )
    r.raise_for_status()
    total = r.json().get("total", 0)

    if total == 0:
        sys.exit("❌ Geen leerlingen gevonden — klopt de manager ID nog?")

    # Stap 2: alles in één keer ophalen 
    params = {"invitedOnAfter": "all-time", "pagesize": total, "sort": "learnerName"}
    r = requests.get(base_url, headers=headers, params=params, timeout=30)
    r.raise_for_status()
    return r.json().get("data", [])


def clean_name(name: str) -> str:
    """Mindtickle zet soms een team-tag als prefix, bv. '[V]Jan Jansen'."""
    name = name or ""
    if name.startswith("[") and "]" in name:
        return name.split("]", 1)[1].strip()
    return name.strip()


def deduplicate_members(members):
    """Houdt per naam alleen het record met de hoogste voltooiing bij."""
    unique_members = {}
    for member in members:
        name = clean_name(member.get("learnerName", ""))
        completion = member.get("percentCompleted", 0) or 0
        existing = unique_members.get(name)
        if existing is None or completion > (existing.get("percentCompleted", 0) or 0):
            unique_members[name] = member
    return list(unique_members.values())


def build_excel(members, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Voortgang"

    headers = [
        "Naam", "Modules", "Voltooiing %", "Modules over tijd",
        "Score %", "Certificeringen",
        f"Boven {SCORE_THRESHOLD:.0f}%?",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border

    # Sorteer op score, hoogste bovenaan
    members_sorted = sorted(
        deduplicate_members(members),
        key=lambda m: m.get("percentScore", 0) or 0,
        reverse=True,
    )

    row_idx = 2
    for m in members_sorted:
        name = clean_name(m.get("learnerName", ""))
        modules = m.get("totalModules", "")
        voltooiing = (m.get("percentCompleted", 0) or 0) / 100
        overdue = m.get("totalOverdue", "")
        score = (m.get("percentScore", 0) or 0) / 100
        cert_pct = m.get("percentCertifications", 0) or 0
        cert_recv = m.get("totalCertificationsReceived", 0) or 0
        cert_total = m.get("totalCertifications", 0) or 0
        cert_str = f"{cert_pct:g}% ({cert_recv}/{cert_total})"

        row_values = [name, modules, voltooiing, overdue, score, cert_str]
        for col, val in enumerate(row_values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.font = Font(name="Arial", size=11)
            cell.border = border
            if col in (3, 5):
                cell.number_format = "0.0%"
            if col != 1:
                cell.alignment = Alignment(horizontal="center")

        status_col = len(row_values) + 1
        status_cell = ws.cell(
            row=row_idx,
            column=status_col,
            value=f'=IF(E{row_idx}>={SCORE_THRESHOLD / 100},"Ja","Nee")',
        )
        status_cell.font = Font(name="Arial", size=11, bold=True)
        status_cell.alignment = Alignment(horizontal="center")
        status_cell.border = border

        row_idx += 1

    last_row = row_idx - 1

    red_fill = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
    green_fill = PatternFill(start_color="C6E0B4", end_color="C6E0B4", fill_type="solid")

    ws.conditional_formatting.add(
        f"E2:E{last_row}",
        CellIsRule(operator="lessThan", formula=[str(SCORE_THRESHOLD / 100)], fill=red_fill),
    )
    ws.conditional_formatting.add(
        f"E2:E{last_row}",
        CellIsRule(operator="greaterThanOrEqual", formula=[str(SCORE_THRESHOLD / 100)], fill=green_fill),
    )
    ws.conditional_formatting.add(
        f"G2:G{last_row}",
        CellIsRule(operator="equal", formula=['"Nee"'], fill=red_fill),
    )
    ws.conditional_formatting.add(
        f"G2:G{last_row}",
        CellIsRule(operator="equal", formula=['"Ja"'], fill=green_fill),
    )

    summary_row = last_row + 2
    ws.cell(row=summary_row, column=1,
            value=f"Aantal boven {SCORE_THRESHOLD:.0f}%:").font = Font(name="Arial", bold=True, size=11)
    ws.cell(row=summary_row, column=5,
            value=f'=COUNTIF(E2:E{last_row},">={SCORE_THRESHOLD / 100}")').font = Font(name="Arial", bold=True, size=11)
    ws.cell(row=summary_row + 1, column=1,
            value="Totaal aantal personen:").font = Font(name="Arial", bold=True, size=11)
    ws.cell(row=summary_row + 1, column=5,
            value=f"=COUNTA(A2:A{last_row})").font = Font(name="Arial", bold=True, size=11)

    widths = {1: 24, 2: 10, 3: 14, 4: 16, 5: 10, 6: 16, 7: 12}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A2"
    wb.save(output_path)


def main():
    print("📡 Data ophalen bij Mindtickle...")
    members = fetch_members()
    print(f"✅ {len(members)} leerlingen opgehaald.")

    today = datetime.now().strftime("%Y-%m-%d")
    output_dir = Path("ouput")
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"voortgang_overzicht_{today}.xlsx"
    build_excel(members, output_path)
    print(f"📊 Excel opgeslagen als: {output_path}")


if __name__ == "__main__":
    main()
